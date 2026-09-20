# 主转换流程(M7-T2/T4/T5):HTML → python-docx → .docx
# 职责:文档初始化(字体/页面/页眉脚)→ 按 section_model 切节 → 逐块渲染 → 落盘(N5:先写临时文件再改名)
import os
import re
import traceback

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import decorations, images
from .element_map import (add_heading, add_list, add_text_paragraph, apply_inline,
                          set_style_font)
from .errors import HtmlParseError, WriteFailedError
from .images import ImageError
from .page_css import PAPER_SIZES_CM, parse_page_config
from .section_model import is_toc_placeholder, split_sections


class _Ctx:
    """渲染上下文:贯穿整个转换过程的共享状态。"""

    def __init__(self, body_font: str, content_width_emu: int, output_dir: str):
        self.body_font = body_font
        self.content_width_emu = content_width_emu
        self.output_dir = output_dir
        self.warnings: list[dict] = []

    # 行内 <img> 与块级 <img> 共用的落图逻辑(F12:失败 → 占位 + warning,不中断)
    def add_picture_to_run(self, run, img_tag) -> None:
        src = img_tag.get('src') or ''
        try:
            stream = images.load_image_bytes(src, self.output_dir)
            width = images.resolve_display_width(img_tag, self.content_width_emu)
            run.add_picture(stream, width=width)  # 只给宽度时按原图比例缩放
        except ImageError as e:
            self.warnings.append({'kind': 'image_failed', 'src': src[:300], 'reason': str(e)})
            run.text = f'[图片加载失败: {e}]'
        except Exception as e:  # python-docx 层面的图片异常(如极端格式)同样只降级不中断
            self.warnings.append({'kind': 'image_failed', 'src': src[:300], 'reason': f'嵌入失败: {e}'})
            run.text = f'[图片加载失败: {e}]'


def convert(html: str, output_path: str) -> dict:
    """入口:HTML 字符串 + 输出路径 → 结果契约 dict(见 __main__ 模块注释)。"""
    soup = BeautifulSoup(html, 'lxml')
    body = soup.body or soup
    if not body.find(True) and not (body.get_text() or '').strip():
        raise HtmlParseError('HTML 中没有可转换的内容')

    config = parse_page_config(soup)

    doc = Document()
    _init_document_fonts(doc, config['body_font'])
    content_width_emu = _apply_page_setup(doc, config)
    _apply_header_footer(doc, config)

    ctx = _Ctx(config['body_font'], content_width_emu,
               os.path.dirname(os.path.abspath(output_path)) or '.')

    sections = split_sections(body)
    for s_idx, section_elements in enumerate(sections):
        # 节首标记:非第一节的首块需要"段前分页"(F5/F14)
        pending_break = s_idx > 0
        for el in _iter_blocks(section_elements):
            pending_break = _render_block(doc, el, ctx, pending_break)

    # 落盘(N5):先写 .tmp,成功后原子改名;失败清理残件
    tmp_path = output_path + '.tmp'
    try:
        parent = os.path.dirname(os.path.abspath(output_path))
        if parent and not os.path.isdir(parent):
            raise WriteFailedError(f'输出目录不存在: {parent}')
        doc.save(tmp_path)
        os.replace(tmp_path, output_path)
    except WriteFailedError:
        _cleanup(tmp_path)
        raise
    except OSError as e:
        _cleanup(tmp_path)
        raise WriteFailedError(f'写入输出文件失败: {e}', detail=output_path) from e

    return {'ok': True, 'outputPath': output_path, 'warnings': ctx.warnings}


# ---- 块级遍历与分发 ----

def _iter_blocks(elements):
    """展平顶层元素:section 容器已由 split_sections 拆开,这里透传列表。"""
    for el in elements:
        yield el


def _render_block(doc, el, ctx, pending_break: bool) -> bool:
    """渲染一个块级元素;返回新的 pending_break 状态。

    分页实现(2026-09-15 实测定案):docx-preview 只按显式分页符(w:br type="page")
    与分节符切页,**不响应 w:pageBreakBefore 属性**(源码解析后未使用)。
    因此节首分页统一为"块前插入独立分页符段落",Word 与预览器行为一致:
      - F14 边界一(表格无法承载段前分页):分页符段落本身就是锚点,天然覆盖;
      - F14 边界二(目录占位被移除):分页符段落独立于占位存在,分页不会随占位丢失。
    pending_break=True 表示当前块是某节的首块,渲染前先落分页符段落。
    """
    name = el.name.lower()

    # ---- 跳过不可见头部内容 ----
    if name in ('style', 'script', 'meta', 'link', 'title', 'head'):
        return pending_break

    # ---- 语义分节容器:递归渲染子块,节首块前落分页符(F5) ----
    if name == 'section':
        for child in el.children:
            if isinstance(child, Tag):
                pending_break = _render_block(doc, child, ctx, pending_break)
            elif isinstance(child, NavigableString) and str(child).strip():
                if pending_break:
                    _break_paragraph(doc)
                    pending_break = False
                paragraph = doc.add_paragraph()
                apply_inline(paragraph, type('_Shim', (), {'children': [child]})(), {}, ctx)
        return pending_break

    # ---- 目录占位(F10) ----
    if is_toc_placeholder(el):
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        _add_toc_field(doc)
        return pending_break

    # ---- 标题 ----
    if name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        add_heading(doc, el, int(name[1]), ctx)
        return pending_break

    # ---- 段落 ----
    if name == 'p':
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        add_text_paragraph(doc, el, ctx)
        return pending_break

    # ---- 列表 ----
    if name in ('ul', 'ol'):
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        add_list(doc, el, ctx)
        return pending_break

    # ---- 表格 ----
    if name == 'table':
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        _render_table(doc, el, ctx)
        return pending_break

    # ---- 块级图片 ----
    if name == 'img':
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        paragraph = doc.add_paragraph()
        ctx.add_picture_to_run(paragraph.add_run(), el)
        return pending_break

    # ---- 分隔线 ----
    if name == 'hr':
        decorations.render_hr(doc, ctx)
        return pending_break  # hr 不承载分页语义,分页符顺延到下一块

    # ---- 装饰组件(F13) ----
    if name == 'div':
        classes = el.get('class') or []
        if 'callout' in classes:
            if pending_break:
                _break_paragraph(doc)
                pending_break = False
            decorations.render_callout(doc, el, ctx)
            return pending_break
        if 'stat-cards' in classes:
            if pending_break:
                _break_paragraph(doc)
                pending_break = False
            decorations.render_stat_cards(doc, el, ctx)
            return pending_break
        # 普通容器:递归渲染子块;纯文本子节点转段落
        for child in el.children:
            if isinstance(child, Tag):
                pending_break = _render_block(doc, child, ctx, pending_break)
            elif isinstance(child, NavigableString) and str(child).strip():
                if pending_break:
                    _break_paragraph(doc)
                    pending_break = False
                paragraph = doc.add_paragraph()
                apply_inline(paragraph, type('_Shim', (), {'children': [child]})(), {}, ctx)
        return pending_break

    # ---- 其他未知标签:含块级子元素则按容器递归;否则按行内段落渲染,不丢内容 ----
    if el.find(['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'ul', 'ol', 'table',
                'section', 'div', 'nav', 'blockquote', 'pre', 'img', 'hr']) is not None:
        for child in el.children:
            if isinstance(child, Tag):
                pending_break = _render_block(doc, child, ctx, pending_break)
            elif isinstance(child, NavigableString) and str(child).strip():
                if pending_break:
                    _break_paragraph(doc)
                    pending_break = False
                paragraph = doc.add_paragraph()
                apply_inline(paragraph, type('_Shim', (), {'children': [child]})(), {}, ctx)
        return pending_break

    if (el.get_text() or '').strip():
        if pending_break:
            _break_paragraph(doc)
            pending_break = False
        paragraph = doc.add_paragraph()
        apply_inline(paragraph, el, {}, ctx)

    return pending_break


def _break_paragraph(doc) -> None:
    """插入独立分页符段落(w:br type="page",Word 与 docx-preview 都按它切页)。

    字号压到 1pt、间距归零,上一页末尾不留可见空行。
    """
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)
    run.font.size = Pt(1)


def _render_table(doc, el, ctx) -> None:
    from .table_style_applier import render_table

    render_table(doc, el, ctx)


# ---- 文档初始化 ----

_HEADING_FONT_CANDIDATES = ('黑体', '微软雅黑', 'Microsoft YaHei')
_HEADING_ACCENT_COLOR = '1F3864'  # 深蓝灰,比纯黑柔和,比模板蓝更稳重


def _init_document_fonts(doc, body_font: str) -> None:
    """文档级字体(F2):Normal 设正文字体;Heading 1-6 换黑体系并配深色强调色,
    与正文形成字体反差(纯黑同字体的标题在视觉上和正文无区别,是最常见的"难看"信号)。
    """
    set_style_font(doc.styles['Normal'], body_font)
    heading_font = _HEADING_FONT_CANDIDATES[0]
    for level in range(1, 7):
        style = doc.styles[f'Heading {level}']
        set_style_font(style, heading_font)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(_HEADING_ACCENT_COLOR)


def _apply_page_setup(doc, config: dict) -> int:
    """纸张与边距(F6),返回版心宽(EMU)供图片钳制。"""
    section = doc.sections[0]
    width_cm, height_cm = PAPER_SIZES_CM[config['paper']]
    section.page_width = Cm(width_cm)
    section.page_height = Cm(height_cm)
    margins = config['margins_cm']
    section.top_margin = Cm(margins['top'])
    section.bottom_margin = Cm(margins['bottom'])
    section.left_margin = Cm(margins['left'])
    section.right_margin = Cm(margins['right'])
    return int(Cm(width_cm) - Cm(margins['left']) - Cm(margins['right']))


def _apply_header_footer(doc, config: dict) -> None:
    """页眉页脚(F11):{{page}}/{{pages}} → PAGE/NUMPAGES 域;首页可单独留空。"""
    section = doc.sections[0]

    if config['header_text']:
        header = section.header
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        _fill_with_fields(paragraph, config['header_text'])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if config['footer_text']:
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        _fill_with_fields(paragraph, config['footer_text'])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 页码惯例居中

    if config['first_page_plain']:
        section.different_first_page_header_footer = True
        # 首页眉脚留空即达成"封面无页眉页脚"(F11)


def _fill_with_fields(paragraph, text: str) -> None:
    """把含 {{page}}/{{pages}} 占位符的文本写进段落,占位符替换为 Word 域。"""
    pattern = re.compile(r'\{\{\s*(page|pages)\s*\}\}', re.IGNORECASE)
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        instr = ' PAGE ' if m.group(1).lower() == 'page' else ' NUMPAGES '
        _add_field(paragraph, instr)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_field(paragraph, instr: str, hint: str | None = None) -> None:
    """插入复杂域(fldChar begin → instrText → separate → 提示文本 → end)。"""
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = instr
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(separate)
    if hint:
        paragraph.add_run(hint)
    end_run = paragraph.add_run()
    end_run._r.append(end)


def _add_toc_field(doc):
    """目录域(F10):收录三级标题、可跳转;begin 带 dirty 让 Word 提示更新页码。"""
    paragraph = doc.add_paragraph()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')  # 打开文档时 Word 会提示"更新域"刷新页码
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = r' TOC \o "1-3" \h \z \u '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(separate)
    hint_run = paragraph.add_run('(目录:在 Word 中右键此处选择「更新域」以生成条目与页码)')
    hint_run.font.color.rgb = RGBColor.from_string('808080')
    end_run = paragraph.add_run()
    end_run._r.append(end)
    return paragraph


def _cleanup(path: str) -> None:
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError:
        pass  # 清理失败不影响错误上报
