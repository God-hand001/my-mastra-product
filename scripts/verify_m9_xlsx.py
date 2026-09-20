# M9 A 组断言脚本
# 用法: vendor/python/python.exe scripts/verify_m9_xlsx.py
# 前置: node scripts/setup-python-runtime.mjs 已执行,openpyxl/formulas/xlsx_recalc 已可用
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'vendor' / '_verify_out' / 'm9'
OUT.mkdir(parents=True, exist_ok=True)
BUNDLE = OUT / 'xlsx-build-bundle.js'
HARNESS = OUT / 'xlsx-build-harness.mjs'
WORKSPACE = ROOT / 'src' / 'mastra' / 'public' / 'workspace'
PY = sys.executable
NODE = shutil.which('node')

CHECKS = []


def check(name, cond, detail=''):
    CHECKS.append((name, bool(cond)))
    print(('PASS' if cond else 'FAIL'), name, ('-- ' + detail if detail and not cond else ''))


def ensure_bundle():
    """把 TypeScript 工具打包成单个 JS,供 Node 端调用。"""
    if BUNDLE.exists() and HARNESS.exists():
        return
    esbuild = ROOT / 'node_modules' / 'esbuild' / 'bin' / 'esbuild'
    subprocess.run(
        [str(NODE), str(esbuild), str(ROOT / 'src' / 'mastra' / 'tools' / 'xlsx-build.ts'),
         '--bundle', '--platform=node', '--format=esm', f'--outfile={BUNDLE}'],
        check=True, cwd=ROOT, timeout=120,
    )
    harness = '''import { xlsxBuildTool } from './xlsx-build-bundle.js';\nprocess.stdin.setEncoding('utf-8');\nlet data = '';\nprocess.stdin.on('data', c => data += c);\nprocess.stdin.on('end', async () => {\n  try {\n    const task = JSON.parse(data || '{}');\n    const r = await xlsxBuildTool.execute(task);\n    console.log(JSON.stringify({ ok: true, result: r }));\n  } catch (e) {\n    console.log(JSON.stringify({ ok: false, error: String(e && e.message ? e.message : e) }));\n  }\n});'''
    HARNESS.write_text(harness, encoding='utf-8')


def call_tool(payload: dict) -> dict:
    ensure_bundle()
    r = subprocess.run(
        [str(NODE), str(HARNESS)],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True, text=True, encoding='utf-8', errors='replace',
        cwd=OUT, timeout=120,
    )
    lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip()]
    if not lines:
        return {'ok': False, 'error': f'无输出: {r.stderr[:400]}'}
    parsed = json.loads(lines[-1])
    if not parsed.get('ok'):
        return {'ok': False, 'error': parsed.get('error', '未知错误')}
    return parsed['result']


def make_sample_xlsx(path: Path, with_error=False, with_external=False, with_dynamic=False):
    """构造测试用 xlsx。"""
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '测试'
    ws['A1'] = '项目'; ws['B1'] = '数值'
    ws['A2'] = 'A'; ws['B2'] = 10
    ws['A3'] = 'B'; ws['B3'] = 20
    ws['A4'] = 'C'; ws['B4'] = 0
    ws['D1'] = '=SUM(B2:B4)'
    ws['D2'] = '=AVERAGE(B2:B4)'
    ws['D3'] = '=IF(B4=0,0,B2/B4)'
    ws['D4'] = '=VLOOKUP("A",A2:B4,2,FALSE)'
    ws['D5'] = '=IF(B2>5,"大","小")'
    ws['E1'] = '=SUMPRODUCT(B2:B4,B2:B4)'
    if with_error:
        ws['E2'] = '=B2/B4'
    if with_external:
        ws['F1'] = '=[other.xlsx]Sheet1!A1'
    if with_dynamic:
        ws['G1'] = ArrayFormula('G1:G3', 'A2:A4*B2:B4')
    wb.save(path)


# ---------- A1 运行时依赖冒烟 ----------
try:
    import openpyxl
    import formulas
    import xlsx_recalc
    check('A1 运行时依赖可用(openpyxl/formulas/xlsx_recalc)', True)
except Exception as e:
    check('A1 运行时依赖可用', False, str(e))
    print('A1 失败,后续断言跳过')
    sys.exit(1)


# ---------- A2 recalc 全对 + backfilled ----------
ok_file = OUT / 'a2-ok.xlsx'
make_sample_xlsx(ok_file)
r = subprocess.run([PY, '-m', 'xlsx_recalc', str(ok_file), '30'],
                   capture_output=True, text=True, encoding='utf-8', errors='replace',
                   env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
res = json.loads(lines[-1]) if lines else {}
check('A2 recalc status=success', res.get('status') == 'success', json.dumps(res, ensure_ascii=False)[:300])
check('A2 totalErrors=0', res.get('totalErrors') == 0)
check('A2 backfilled=true', res.get('backfilled') is True)
wb = openpyxl.load_workbook(str(ok_file), data_only=True)
ws = wb['测试']
cached = [ws['D1'].value, ws['D2'].value, ws['D3'].value, ws['D4'].value, ws['D5'].value, ws['E1'].value]
check('A2 公式缓存值非空', all(v is not None for v in cached), str(cached))


# ---------- A3 errors_found 定位 + 退出码 0 ----------
div_file = OUT / 'a3-div0.xlsx'
make_sample_xlsx(div_file, with_error=True)
r = subprocess.run([PY, '-m', 'xlsx_recalc', str(div_file), '30'],
                   capture_output=True, text=True, encoding='utf-8', errors='replace',
                   env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
res = json.loads(lines[-1]) if lines else {}
check('A3 recalc status=errors_found', res.get('status') == 'errors_found', json.dumps(res, ensure_ascii=False)[:300])
check('A3 totalErrors>=1', (res.get('totalErrors') or 0) >= 1)
summary = res.get('errorSummary') or []
check('A3 errorSummary 含 div0', any(s.get('kind') == 'div0' for s in summary), str(summary))
check('A3 errors_found 退出码为 0', r.returncode == 0, f'rc={r.returncode}')


# ---------- A4 外部引用与动态数组 = error ----------
for name, kw in [('external', {'with_external': True}), ('dynamic', {'with_dynamic': True})]:
    fp = OUT / f'a4-{name}.xlsx'
    make_sample_xlsx(fp, **kw)
    r = subprocess.run([PY, '-m', 'xlsx_recalc', str(fp), '30'],
                       capture_output=True, text=True, encoding='utf-8', errors='replace',
                       env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
    lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
    res = json.loads(lines[-1]) if lines else {}
    check(f'A4 {name} status=error', res.get('status') == 'error', json.dumps(res, ensure_ascii=False)[:300])


# ---------- A5 工具闭环 ----------
tool_script = '''
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = '预算'
ws['A1'] = '项目'
ws['B1'] = '金额'
ws['A2'] = '房租'
ws['B2'] = 3000
ws['A3'] = '餐饮'
ws['B3'] = 1500
ws['A4'] = '合计'
ws['B4'] = '=SUM(B2:B3)'
ws['B5'] = '=IF(B3=0,0,B2/B3)'

header_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')
for col in ['A','B']:
    ws.column_dimensions[col].width = 16
wb.save(sys.argv[1])
'''
rel_path = 'm9_verify/月度开支.xlsx'
res = call_tool({'buildScript': tool_script, 'xlsxPath': rel_path})
check('A5 工具闭环 ok', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:400])
if res.get('ok'):
    abs_path = WORKSPACE / rel_path
    check('A5 产物文件存在', abs_path.exists())
    if abs_path.exists():
        wb = openpyxl.load_workbook(str(abs_path), data_only=True)
        check('A5 合计缓存值正确', wb['预算']['B4'].value == 4500, str(wb['预算']['B4'].value))


# ---------- A6 pip install 拒绝 ----------
bad_script = '''
import subprocess
subprocess.run(['pip', 'install', 'openpyxl'])
'''
res = call_tool({'buildScript': bad_script, 'xlsxPath': 'm9_verify/bad.xlsx'})
check('A6 pip install 被拒绝', res.get('ok') is False and 'pip install' in (res.get('error') or ''),
      json.dumps(res, ensure_ascii=False)[:400])


# ---------- A7 语法错误 stderr 定位 ----------
syntax_script = '''
from openpyxl import Workbook
wb = Workbook(
wb.save(sys.argv[1])
'''
res = call_tool({'buildScript': syntax_script, 'xlsxPath': 'm9_verify/syntax.xlsx'})
check('A7 语法错误返回失败', res.get('ok') is False)
check('A7 错误含行号/语法信息', res.get('error') and ('SyntaxError' in res['error'] or 'line' in res['error'].lower()),
      res.get('error'))


# ---------- A8 clean 清理 ----------
clean_rel = 'm9_verify/clean_test.xlsx'
clean_res = call_tool({'buildScript': tool_script, 'xlsxPath': clean_rel})
base = Path(clean_rel).stem
ref_dir = WORKSPACE / f'.{base}.ref'
check('A8 clean 前 .ref 存在', ref_dir.exists())
clean_res2 = call_tool({'buildScript': '', 'xlsxPath': clean_rel, 'clean': True})
check('A8 clean 返回 ok+cleaned', clean_res2.get('ok') is True and clean_res2.get('cleaned') is True)
check('A8 clean 后 .ref 消失', not ref_dir.exists())


# ---------- A9 防穿越 ----------
res = call_tool({'buildScript': tool_script, 'xlsxPath': '../etc.xlsx'})
check('A9 非法路径被拒绝', res.get('ok') is False and '非法' in (res.get('error') or ''),
      json.dumps(res, ensure_ascii=False)[:400])


# ---------- A10 技能文件齐 5 个 + frontmatter 可解析 ----------
skill_dir = ROOT / 'extensions' / 'skills' / 'excel-generation'
files = [
    skill_dir / 'SKILL.md',
    skill_dir / 'references' / 'schema_principle.md',
    skill_dir / 'references' / 'patterns' / '预算表.md',
    skill_dir / 'references' / 'patterns' / '进度表.md',
    skill_dir / 'references' / 'patterns' / '统计卡表.md',
]
check('A10 技能文件齐 5 个', all(f.exists() for f in files), str([str(f) for f in files if not f.exists()]))
skill_md = (skill_dir / 'SKILL.md').read_text(encoding='utf-8')
fm = skill_md.split('---', 2)
if len(fm) >= 3:
    meta_lines = [ln for ln in fm[1].strip().splitlines() if ':' in ln]
    meta = {}
    for ln in meta_lines:
        k, v = ln.split(':', 1)
        meta[k.strip()] = v.strip().strip('"')
    check('A10 SKILL.md frontmatter 可解析', 'name' in meta and 'description' in meta, str(meta))
else:
    check('A10 SKILL.md frontmatter 可解析', False, '未找到 frontmatter')


# ---------- A11 M7 转换器回归 ----------
m7 = subprocess.run([PY, 'scripts/verify_m7_converter.py'], cwd=ROOT,
                    capture_output=True, text=True, encoding='utf-8', errors='replace',
                    env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
check('A11 verify_m7_converter.py 仍绿', m7.returncode == 0, m7.stdout[-600:] + m7.stderr[-300:])


# ---------- 汇总 ----------
failed = [n for n, ok in CHECKS if not ok]
print(f"\n总计 {len(CHECKS)} 项,通过 {len(CHECKS) - len(failed)},失败 {len(failed)}")
if failed:
    print('失败项:', ', '.join(failed))
    sys.exit(1)
print('全部通过 ✅')
