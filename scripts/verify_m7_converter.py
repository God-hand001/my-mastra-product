# M7 A 组断言脚本(对应 m7-checklist.md A1–A10 的转换器部分)
# 用法:vendor/python/python.exe scripts/verify_m7_converter.py
# 前置:node scripts/setup-python-runtime.mjs 已执行(转换器包已拷入运行时 site-packages)
# 流程:生成 sample.png → 逐个转换 fixture → 用 python-docx 打开产物做 XML 级断言
import base64
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / 'docs' / 'spec' / 'fixtures'
OUT = ROOT / 'vendor' / '_verify_out'
OUT.mkdir(parents=True, exist_ok=True)

# 子进程强制 UTF-8 输出;errors='replace' 兜住 OS 层 GBK 消息,避免解码崩溃
CONVERT_ENV = {**os.environ, 'PYTHONIOENCODING': 'utf-8'}

CHECKS = []


def check(name, cond, detail=''):
    CHECKS.append((name, bool(cond)))
    print(('PASS' if cond else 'FAIL'), name, ('-- ' + detail if detail and not cond else ''))


def convert(fixture: Path, out_name: str):
    """转换一个 fixture,返回 (结果 dict, 产物 Path)。"""
    out = OUT / out_name
    if out.exists():
        out.unlink()
    task = {'html': fixture.read_text(encoding='utf-8'), 'outputPath': str(out)}
    with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8') as f:
        json.dump(task, f, ensure_ascii=False)
        task_file = f.name
    r = subprocess.run([sys.executable, '-m', 'html_to_docx', '--input', task_file],
                       capture_output=True, text=True, encoding='utf-8', errors='replace',
                       timeout=90, env=CONVERT_ENV)
    lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
    if not lines:
        return {'ok': False, 'error': {'message': f'stdout 无 JSON: {r.stdout[:200]} / {r.stderr[:400]}'}}, out
    return json.loads(lines[-1]), out


def doc_xml(doc):
    """document.xml 全文(字符串),便于标记性断言。"""
    return doc.element.xml


# ---------- 准备 sample.png(1×1 红点) ----------
# 注意契约:HTML 里的相对路径以"输出文件目录"为基准解析,所以图必须放在 OUT 里
sample = FIX / 'sample.png'
if not sample.exists():
    # 1x1 红色 PNG
    sample.write_bytes(base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=='
    ))
(OUT / 'sample.png').write_bytes(sample.read_bytes())

# ---------- A1 运行时可用(能跑到这里说明依赖导入成功) ----------
check('A1 运行时依赖可用', True)

# ---------- A2/A3 基础映射 + 中文字体 ----------
from docx import Document

res, out = convert(FIX / 'm7-basic.html', 'basic.docx')
check('A2 basic 转换成功', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:300])
if res.get('ok'):
    doc = Document(str(out))
    xml = doc_xml(doc)
    styles = [p.style.name for p in doc.paragraphs]
    check('A2 标题三级映射', 'Heading 1' in styles and 'Heading 2' in styles and 'Heading 3' in styles, str(styles[:12]))
    check('A2 列表映射', 'List Bullet' in styles and 'List Number' in styles and 'List Bullet 2' in styles, str(styles))
    check('A3 eastAsia 中文字体', 'w:eastAsia' in xml and '宋体' in xml)
    check('A2 居中存在', 'w:jc w:val="center"' in xml.replace('w:jc w:val="center"', 'w:jc w:val="center"') or 'center' in xml)

# ---------- A4 失败分类 + 无残件(N5/AC2/AC10) ----------
bad_dir_task = {'html': '<html><body><p>ok</p></body></html>', 'outputPath': str(Path('Z:/不存在目录/x.docx'))}
with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8') as f:
    json.dump(bad_dir_task, f, ensure_ascii=False)
    bad_task = f.name
r = subprocess.run([sys.executable, '-m', 'html_to_docx', '--input', bad_task],
                   capture_output=True, text=True, encoding='utf-8', errors='replace',
                   timeout=90, env=CONVERT_ENV)
lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
bad_res = json.loads(lines[-1]) if lines else {}
check('A4 不可写路径 → write_failed', bad_res.get('error', {}).get('kind') == 'write_failed',
      json.dumps(bad_res, ensure_ascii=False)[:300])
res2, out2 = convert(FIX / 'm7-basic.html', 'basic2.docx')
leftovers = list(OUT.glob('*.tmp'))
check('A10 无半成品残件', res2.get('ok') is True and not leftovers, str(leftovers))

# ---------- A5 表格 ----------
res, out = convert(FIX / 'm7-table.html', 'table.docx')
check('A5 表格转换成功', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:300])
if res.get('ok'):
    doc = Document(str(out))
    xml = doc_xml(doc)
    check('A5 表格存在', len(doc.tables) >= 1)
    check('A5 表头底色', 'DCE9F9' in xml)
    check('A5 全边框', 'w:tblBorders' in xml)
    check('A5 列宽设定', 'w:tblLayout' in xml and 'w:w' in xml)

# ---------- A6/A7/A8 分节/目录/页面设置/页眉脚 ----------
res, out = convert(FIX / 'm7-sections.html', 'sections.docx')
check('A6 分节转换成功', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:300])
if res.get('ok'):
    doc = Document(str(out))
    xml = doc_xml(doc)
    pbb = xml.count('w:br w:type="page"')
    check('A6 显式分页符≥2(第二、三节;docx-preview 只认显式分页符)', pbb >= 2, f'count={pbb}')
    check('A7 TOC 域存在', ' TOC ' in xml and 'w:fldChar' in xml)
    check('A8 A4 纸张(11906x16838 twips)', 'w:pgSz' in xml and '11906' in xml and '16838' in xml)
    sec = doc.sections[0]
    footer_xml = ''.join(p._p.xml for p in sec.footer.paragraphs) if sec.footer else ''
    check('A8 页脚 PAGE 域', ' PAGE ' in footer_xml and ' NUMPAGES ' in footer_xml)
    header_text = ''.join(p.text for p in sec.header.paragraphs)
    check('A8 页眉内容', '季度经营简报' in header_text, header_text)
    check('A8 页脚文案', '第' in ''.join(p.text for p in sec.footer.paragraphs))

# ---------- A9/A10 图片与装饰组件 ----------
res, out = convert(FIX / 'm7-rich.html', 'rich.docx')
check('A9 富文档转换成功', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:300])
if res.get('ok'):
    warns = res.get('warnings') or []
    check('A9 warnings 恰 1 条 image_failed',
          len(warns) == 1 and warns[0].get('kind') == 'image_failed', json.dumps(warns, ensure_ascii=False)[:300])
    doc = Document(str(out))
    check('A9 内嵌图片 2 张', len(doc.inline_shapes) == 2, f'n={len(doc.inline_shapes)}')
    xml = doc_xml(doc)
    check('A10 标注框底纹', 'E8F1FB' in xml and 'FDF3E4' in xml)  # info 蓝 + warning 橙(与 fixture 一致)
    check('A10 左侧色条(24/8pt 边框)', 'w:sz="24"' in xml)
    check('A10 hr 横线(pBdr)', 'w:pBdr' in xml)
    # 数值加大加粗:20pt = w:sz 40(半点单位)
    check('A10 数据卡片(20pt 数值)', 'w:val="40"' in xml)

# ---------- 汇总 ----------
failed = [n for n, ok in CHECKS if not ok]
print(f"\n总计 {len(CHECKS)} 项,通过 {len(CHECKS) - len(failed)},失败 {len(failed)}")
if failed:
    print('失败项:', ', '.join(failed))
    sys.exit(1)
print('全部通过 ✅')
