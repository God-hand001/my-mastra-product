# M8 A 组断言脚本(对应 m8-checklist.md A1–A11 的转换器与工具部分)
# 用法:vendor/python/python.exe scripts/verify_m8_pptx.py
# 前置:node scripts/setup-python-runtime.mjs 已执行(转换器包已拷入运行时 site-packages)
import base64
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / 'docs' / 'spec' / 'fixtures'
OUT = ROOT / 'vendor' / '_verify_out'
OUT.mkdir(parents=True, exist_ok=True)

CONVERT_ENV = {**os.environ, 'PYTHONIOENCODING': 'utf-8'}

CHECKS = []


def check(name, cond, detail=''):
    CHECKS.append((name, bool(cond)))
    print(('PASS' if cond else 'FAIL'), name, ('-- ' + detail if detail and not cond else ''))


def convert(fixture: Path, out_name: str):
    """转换一个 fixture,返回 (结果 dict, 产物 Path)。"""
    out = OUT / out_name
    if out.exists():
        out.unlink()
    task = {'html': fixture.read_text(encoding='utf-8'), 'outputPath': str(out)}
    with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8') as f:
        json.dump(task, f, ensure_ascii=False)
        task_file = f.name
    r = subprocess.run([sys.executable, '-m', 'html_to_pptx', '--input', task_file],
                       capture_output=True, text=True, encoding='utf-8', errors='replace',
                       timeout=90, env=CONVERT_ENV)
    lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
    if not lines:
        return {'ok': False, 'error': {'message': f'stdout 无 JSON: {r.stdout[:200]} / {r.stderr[:400]}'}}, out
    return json.loads(lines[-1]), out


# ---------- A1 运行时依赖可用 ----------
from pptx import Presentation

check('A1 运行时依赖可用(pptx 可导入)', True)

# ---------- A2 页数 ----------
res, out = convert(FIX / 'm8-deck.html', 'm8-deck.pptx')
check('A2 m8-deck 转换成功', res.get('ok') is True, json.dumps(res, ensure_ascii=False)[:300])
if res.get('ok'):
    prs = Presentation(str(out))
    check('A2 页数为 4', len(prs.slides) == 4, f'n={len(prs.slides)}')

# ---------- A3 文本 ----------
if res.get('ok'):
    prs = Presentation(str(out))
    all_text = '\n'.join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    check('A3 包含封面标题', '嘉立创办公' in all_text, all_text[:200])
    check('A3 包含内容页标题', '核心能力' in all_text)
    check('A3 包含两栏页标题', '服务版图' in all_text)
    check('A3 包含数据页标题', '经营数据' in all_text)

# ---------- A4 eastAsia 中文字体 ----------
if res.get('ok'):
    prs = Presentation(str(out))
    run_xmls = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run_xmls.append(run._r.xml)
    all_run_xml = ''.join(run_xmls)
    check('A4 XML 含 a:ea 节点', '<a:ea' in all_run_xml)
    check('A4 eastAsia typeface 为微软雅黑', 'typeface="微软雅黑"' in all_run_xml)
    check('A4 XML 含 a:latin 节点', '<a:latin' in all_run_xml)

# ---------- A5 版式坐标 ----------
if res.get('ok'):
    prs = Presentation(str(out))
    # 第 2 页(内容页)标题区应在左上角
    title_shape = [sh for sh in prs.slides[1].shapes if sh.has_text_frame and '核心能力' in sh.text_frame.text]
    if title_shape:
        sh = title_shape[0]
        check('A5 内容页标题区 x≈0.5in', abs(sh.left.inches - 0.5) < 0.1, f'x={sh.left.inches}')
        check('A5 内容页标题区 y≈0.5in', abs(sh.top.inches - 0.5) < 0.1, f'y={sh.top.inches}')
    # 第 3 页(两栏页)左右栏应有多个 shape
    check('A5 两栏页 shape 数≥3', len(prs.slides[2].shapes) >= 3, f'n={len(prs.slides[2].shapes)}')
    left_shapes = [sh for sh in prs.slides[2].shapes if sh.left.inches < 4]
    right_shapes = [sh for sh in prs.slides[2].shapes if sh.left.inches > 6]
    check('A5 两栏页左右栏均有内容', len(left_shapes) >= 1 and len(right_shapes) >= 1)

# ---------- A6 图片与 warnings ----------
if res.get('ok'):
    prs = Presentation(str(out))
    pictures = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
    check('A6 内嵌图片 2 张', len(pictures) == 2, f'n={len(pictures)}')
    warns = res.get('warnings') or []
    check('A6 warnings 恰 1 条 image_failed',
          len(warns) == 1 and warns[0].get('kind') == 'image_failed',
          json.dumps(warns, ensure_ascii=False)[:300])

# ---------- A7 装饰组件 ----------
if res.get('ok'):
    prs = Presentation(str(out))
    slide3_xml = prs.slides[3].part.blob.decode('utf-8', errors='ignore')
    # 数据页有 3 个圆角矩形(卡片)
    rounded = [sh for sh in prs.slides[3].shapes if sh.shape_type == 1]
    check('A7 数据卡片形状存在', len(rounded) >= 3, f'n={len(rounded)}')
    # 标注框 warning 背景色 FDF3E4 + 左侧色条 F5A623
    check('A7 标注框 warning 背景色', 'FDF3E4' in slide3_xml)
    check('A7 标注框 warning 强调色', 'F5A623' in slide3_xml)
    # hr 横线:极矮矩形
    thin_lines = [sh for sh in prs.slides[3].shapes
                  if sh.shape_type == 1 and sh.height.inches < 0.05]
    check('A7 hr 横线形状存在', len(thin_lines) >= 1, f'n={len(thin_lines)}')

# ---------- A8 工具文件 ----------
tool_file = ROOT / 'src/mastra/tools/html-to-pptx.ts'
tool_text = tool_file.read_text(encoding='utf-8')
check('A8 工具文件存在', tool_file.exists())
check('A8 工具 id 为 html_to_pptx', "id: 'html_to_pptx'" in tool_text)
check('A8 工具 description 含 HTML 骨架示例', 'MANDATORY HTML skeleton' in tool_text)
check('A8 工具 description 含 data-layout', 'data-layout' in tool_text)

# ---------- A9 防穿越(TS 工具层) ----------
check('A9 工具实现 safeOutputPath', 'safeOutputPath' in tool_text)
check('A9 非法路径返回中文提示', '必须位于 workspace 根目录内' in tool_text)

# ---------- A11 失败无残件 ----------
bad_dir_task = {'html': '<html><body><section><h1>ok</h1></section></body></html>',
                'outputPath': str(Path('Z:/不存在目录/out.pptx'))}
with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8') as f:
    json.dump(bad_dir_task, f, ensure_ascii=False)
    bad_task = f.name
r = subprocess.run([sys.executable, '-m', 'html_to_pptx', '--input', bad_task],
                   capture_output=True, text=True, encoding='utf-8', errors='replace',
                   timeout=90, env=CONVERT_ENV)
lines = [ln for ln in r.stdout.strip().split('\n') if ln.strip().startswith('{')]
bad_res = json.loads(lines[-1]) if lines else {}
check('A11 不可写路径 → write_failed', bad_res.get('error', {}).get('kind') == 'write_failed',
      json.dumps(bad_res, ensure_ascii=False)[:300])
res2, out2 = convert(FIX / 'm8-deck.html', 'm8-deck-no-leftover.pptx')
leftovers = list(OUT.glob('*.tmp'))
check('A11 无半成品残件', res2.get('ok') is True and not leftovers, str(leftovers))

# ---------- 汇总 ----------
failed = [n for n, ok in CHECKS if not ok]
print(f"\n总计 {len(CHECKS)} 项,通过 {len(CHECKS) - len(failed)},失败 {len(failed)}")
if failed:
    print('失败项:', ', '.join(failed))
    sys.exit(1)
print('全部通过')
